# -*- coding: utf-8 -*-
"""
Pruebas unitarias para el detector heurístico de desplazamientos.
"""

import openpyxl
from core.shift_detector import detectar_desplazamiento


def test_shift_detector_sin_desplazamiento():
    wb_p = openpyxl.Workbook()
    ws_p = wb_p.active

    wb_e = openpyxl.Workbook()
    ws_e = wb_e.active

    for r in range(1, 15):
        ws_p.cell(row=r, column=1, value=f"Dato {r}")
        ws_e.cell(row=r, column=1, value=f"Dato {r}")

    offset, conf, msg = detectar_desplazamiento(ws_p, ws_e)
    assert offset is None
    assert msg == ""


def test_shift_detector_fila_desplazada():
    wb_p = openpyxl.Workbook()
    ws_p = wb_p.active

    wb_e = openpyxl.Workbook()
    ws_e = wb_e.active

    # Plantilla en filas 1..10
    for r in range(1, 11):
        ws_p.cell(row=r, column=1, value=f"Dato {r}")

    # Estudiante insertó una fila vacía arriba (sus datos empiezan en fila 2)
    for r in range(1, 11):
        ws_e.cell(row=r + 1, column=1, value=f"Dato {r}")

    offset, conf, msg = detectar_desplazamiento(ws_p, ws_e)
    assert offset == (1, 0)
    assert conf >= 0.70
    assert "+1 fila(s) abajo" in msg
