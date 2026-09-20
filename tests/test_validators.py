# -*- coding: utf-8 -*-
"""
Pruebas unitarias para el catálogo de validadores del Modo B.
"""

import pytest
import openpyxl
from core.models import CriterioRubrica
from core.validators.value_validator import ValueValidator
from core.validators.formula_validator import FormulaValidator
from core.validators.format_validator import FormatValidator
from core.validators.style_validator import StyleValidator
from core.validators.structure_validator import StructureValidator


def test_value_validator_numeric():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = 1500.04

    val = ValueValidator()

    # Coincidencia con tolerancia 0.05
    crit_ok = CriterioRubrica(
        id="C1",
        criterio="Total 1500",
        hoja="Sheet",
        rango="A1",
        tipo_validacion="valor_numerico",
        parametros={"valor": 1500.0, "tolerancia": 0.05},
        puntos=10
    )
    res_ok = val.validar(ws, crit_ok)
    assert res_ok.aprobado is True
    assert res_ok.puntos_obtenidos == 10

    # Falla con tolerancia estricta
    crit_fail = CriterioRubrica(
        id="C2",
        criterio="Total 1500 estricto",
        hoja="Sheet",
        rango="A1",
        tipo_validacion="valor_numerico",
        parametros={"valor": 1500.0, "tolerancia": 0.01},
        puntos=10
    )
    res_fail = val.validar(ws, crit_fail)
    assert res_fail.aprobado is False
    assert res_fail.puntos_obtenidos == 0


def test_value_validator_text_normalized():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["B2"] = "  MÉXICO  "

    val = ValueValidator()
    crit = CriterioRubrica(
        id="C3",
        criterio="País México",
        hoja="Sheet",
        rango="B2",
        tipo_validacion="valor_exacto",
        parametros={"valor": "mexico"},
        puntos=15
    )
    res = val.validar(ws, crit)
    assert res.aprobado is True
    assert res.puntos_obtenidos == 15


def test_formula_validator_bilingual_and_refs():
    wb = openpyxl.Workbook()
    ws = wb.active
    # En openpyxl las fórmulas pueden estar escritas como SUM(C2:C10) o =SUMA(C2:C10)
    ws["D20"] = "=SUMA(C2:C10)"

    val = FormulaValidator()

    # Validar que reconoce SUM o SUMA
    crit = CriterioRubrica(
        id="C4",
        criterio="Suma de columna C",
        hoja="Sheet",
        rango="D20",
        tipo_validacion="formula_flexible",
        parametros={
            "funciones": ["SUM"],
            "referencias": ["C2:C10"]
        },
        puntos=20
    )
    res = val.validar(ws, crit)
    assert res.aprobado is True
    assert res.puntos_obtenidos == 20


def test_format_validator():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["E5"] = 1250.5
    ws["E5"].number_format = '$#,##0.00'

    ws["F5"] = 0.15
    ws["F5"].number_format = '0.00%'

    val = FormatValidator()

    # Moneda
    crit_moneda = CriterioRubrica(
        id="C5",
        criterio="Moneda en E5",
        hoja="Sheet",
        rango="E5",
        tipo_validacion="formato_numero",
        parametros={"formato": "moneda"},
        puntos=10
    )
    assert val.validar(ws, crit_moneda).aprobado is True

    # Porcentaje
    crit_pct = CriterioRubrica(
        id="C6",
        criterio="Porcentaje en F5",
        hoja="Sheet",
        rango="F5",
        tipo_validacion="formato_numero",
        parametros={"formato": "porcentaje"},
        puntos=10
    )
    assert val.validar(ws, crit_pct).aprobado is True


def test_style_validator():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Encabezado"
    ws["A1"].font = openpyxl.styles.Font(bold=True)
    ws["A1"].fill = openpyxl.styles.PatternFill(start_color="FF0056B3", end_color="FF0056B3", fill_type="solid")

    val = StyleValidator()
    crit = CriterioRubrica(
        id="C7",
        criterio="Encabezado negrita y relleno",
        hoja="Sheet",
        rango="A1",
        tipo_validacion="estilo_visual_flexible",
        parametros={"requiere_negrita": True, "requiere_relleno": True},
        puntos=10
    )
    assert val.validar(ws, crit).aprobado is True


def test_structure_validator_merged():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.merge_cells("A1:D1")

    val = StructureValidator()
    crit = CriterioRubrica(
        id="C8",
        criterio="Título combinado A1:D1",
        hoja="Sheet",
        rango="A1:D1",
        tipo_validacion="celdas_combinadas",
        parametros={},
        puntos=10
    )
    assert val.validar(ws, crit).aprobado is True
