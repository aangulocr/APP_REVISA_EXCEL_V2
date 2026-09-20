# -*- coding: utf-8 -*-
# ============================================================================
# core/rubric_template.py — Generador de plantilla Excel con _RUBRICA prellenada
# ============================================================================
"""
Crea un libro de Excel modelo que contiene ejemplos funcionales de datos,
fórmulas, formatos y una hoja '_RUBRICA' completamente configurada que
suma 100 puntos y cubre todos los validadores estándar.
"""

import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from core.config import (
    PLANTILLAS_DIR,
    HOJA_RUBRICA_DEFAULT,
    COLOR_HEADER_BG,
    COLOR_HEADER_FG,
)


def generar_plantilla_rubrica(ruta_salida: str = None) -> str:
    """
    Genera el archivo PLANTILLA_RUBRICA_EJEMPLO.xlsx con hojas 'Ventas', 'Resumen'
    y la hoja '_RUBRICA' con 100 puntos distribuidos en 8 criterios prácticos.
    """
    if not ruta_salida:
        os.makedirs(PLANTILLAS_DIR, exist_ok=True)
        ruta_salida = os.path.join(PLANTILLAS_DIR, "PLANTILLA_RUBRICA_EJEMPLO.xlsx")

    wb = openpyxl.Workbook()

    # Estilos comunes
    font_header = Font(name="Calibri", size=11, bold=True, color=COLOR_HEADER_FG[2:])
    fill_header = PatternFill(start_color=COLOR_HEADER_BG, end_color=COLOR_HEADER_BG, fill_type="solid")
    fill_rubrica_header = PatternFill(start_color="FF2C5282", end_color="FF2C5282", fill_type="solid")
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")

    thin_border = Border(
        left=Side(style='thin', color='D0D7DE'),
        right=Side(style='thin', color='D0D7DE'),
        top=Side(style='thin', color='D0D7DE'),
        bottom=Side(style='thin', color='D0D7DE')
    )

    # -------------------------------------------------------------------------
    # 1. HOJA: Ventas
    # -------------------------------------------------------------------------
    ws_ventas = wb.active
    ws_ventas.title = "Ventas"

    encabezados_ventas = ["Producto", "Categoria", "Cantidad", "Precio_Unitario", "Subtotal", "Descuento", "Total"]
    for col_idx, h in enumerate(encabezados_ventas, 1):
        c = ws_ventas.cell(row=1, column=col_idx, value=h)
        c.font = font_header
        c.fill = fill_header
        c.alignment = align_center
        c.border = thin_border
    ws_ventas.row_dimensions[1].height = 25

    datos_ventas = [
        ("Teclado Mecánico", "Periféricos", 15, 45.0),
        ("Mouse Gamer", "Periféricos", 20, 25.0),
        ("Monitor 27 Pulgadas", "Monitores", 8, 220.0),
        ("Auriculares Inalámbricos", "Audio", 12, 60.0),
        ("Webcam 1080p", "Periféricos", 10, 35.0),
        ("Micrófono USB", "Audio", 5, 80.0),
        ("Soporte Monitor", "Accesorios", 7, 30.0),
        ("Cable HDMI 4K", "Accesorios", 25, 10.0),
        ("Parlantes Bluetooth", "Audio", 9, 45.0),
    ]

    for idx, (prod, cat, cant, precio) in enumerate(datos_ventas, 2):
        ws_ventas.cell(row=idx, column=1, value=prod).border = thin_border
        ws_ventas.cell(row=idx, column=2, value=cat).border = thin_border

        c_cant = ws_ventas.cell(row=idx, column=3, value=cant)
        c_cant.border = thin_border
        c_cant.alignment = align_right
        c_cant.number_format = '#,##0'

        c_prec = ws_ventas.cell(row=idx, column=4, value=precio)
        c_prec.border = thin_border
        c_prec.alignment = align_right
        c_prec.number_format = '$#,##0.00'

        # Fórmulas
        c_sub = ws_ventas.cell(row=idx, column=5, value=f"=C{idx}*D{idx}")
        c_sub.border = thin_border
        c_sub.alignment = align_right
        c_sub.number_format = '$#,##0.00'

        c_desc = ws_ventas.cell(row=idx, column=6, value=f"=E{idx}*0.05")
        c_desc.border = thin_border
        c_desc.alignment = align_right
        c_desc.number_format = '$#,##0.00'

        c_tot = ws_ventas.cell(row=idx, column=7, value=f"=E{idx}-F{idx}")
        c_tot.border = thin_border
        c_tot.alignment = align_right
        c_tot.number_format = '$#,##0.00'

    # Fila de Total General (Fila 11)
    fila_total = len(datos_ventas) + 2
    ws_ventas.cell(row=fila_total, column=1, value="TOTAL GENERAL").font = Font(bold=True)
    for c in range(1, 7):
        ws_ventas.cell(row=fila_total, column=c).border = thin_border

    c_grantotal = ws_ventas.cell(row=fila_total, column=7, value=f"=SUM(G2:G{fila_total-1})")
    c_grantotal.font = Font(bold=True)
    c_grantotal.border = thin_border
    c_grantotal.alignment = align_right
    c_grantotal.number_format = '$#,##0.00'

    # -------------------------------------------------------------------------
    # 2. HOJA: Resumen
    # -------------------------------------------------------------------------
    ws_resumen = wb.create_sheet(title="Resumen")

    # Título Combinado A1:D1
    ws_resumen.merge_cells("A1:D1")
    c_tit = ws_resumen.cell(row=1, column=1, value="PANEL DE CONTROL Y RESUMEN DE VENTAS")
    c_tit.font = Font(name="Calibri", size=13, bold=True, color="FFFFFF")
    c_tit.fill = fill_header
    c_tit.alignment = align_center
    ws_resumen.row_dimensions[1].height = 32

    # Validación de datos en B3
    ws_resumen.cell(row=3, column=1, value="Seleccionar Categoría:").font = Font(bold=True)
    c_val = ws_resumen.cell(row=3, column=2, value="Periféricos")
    c_val.border = thin_border

    dv = DataValidation(type="list", formula1='"Periféricos,Monitores,Audio,Accesorios"', allow_blank=True)
    ws_resumen.add_data_validation(dv)
    dv.add("B3")

    ws_resumen.cell(row=5, column=1, value="Métricas Clave:").font = Font(bold=True)
    ws_resumen.cell(row=6, column=1, value="Total Productos Evaluados:")
    ws_resumen.cell(row=6, column=2, value="=COUNTA(Ventas!A2:A10)")
    ws_resumen.cell(row=7, column=1, value="Promedio de Venta por Ítem:")
    ws_resumen.cell(row=7, column=2, value="=AVERAGE(Ventas!G2:G10)")
    ws_resumen.cell(row=7, column=2).number_format = '$#,##0.00'

    # -------------------------------------------------------------------------
    # 3. HOJA: _RUBRICA
    # -------------------------------------------------------------------------
    ws_rubrica = wb.create_sheet(title=HOJA_RUBRICA_DEFAULT)

    encabezados_rubrica = [
        "ID", "CRITERIO", "HOJA", "RANGO", "TIPO_VALIDACION", "PARAMETROS", "PUNTOS", "OBLIGATORIO"
    ]
    for col_idx, h in enumerate(encabezados_rubrica, 1):
        c = ws_rubrica.cell(row=1, column=col_idx, value=h)
        c.font = font_header
        c.fill = fill_rubrica_header
        c.alignment = align_center
        c.border = thin_border
    ws_rubrica.row_dimensions[1].height = 28

    criterios_ejemplo = [
        (
            "CRIT_01",
            "Calcular Subtotal multiplicando Cantidad por Precio Unitario",
            "Ventas",
            "E2:E10",
            "formula_flexible",
            '{"referencias": ["C", "D"]}',
            15,
            "SI"
        ),
        (
            "CRIT_02",
            "Calcular Total general de Ventas mediante función SUMA",
            "Ventas",
            "G11",
            "formula_flexible",
            '{"funciones": ["SUMA", "SUM"], "referencias": ["G2:G10"]}',
            20,
            "SI"
        ),
        (
            "CRIT_03",
            "Verificar valor numérico del Gran Total de Ventas",
            "Ventas",
            "G11",
            "valor_numerico",
            '{"valor": 5006.5, "tolerancia": 1.0}',
            15,
            "NO"
        ),
        (
            "CRIT_04",
            "Aplicar formato de Moneda en columnas de precio y totales",
            "Ventas",
            "D2:G11",
            "formato_numero",
            '{"formato": "moneda"}',
            10,
            "NO"
        ),
        (
            "CRIT_05",
            "Aplicar estilo visual a encabezados (negrita, relleno y bordes)",
            "Ventas",
            "A1:G1",
            "estilo_visual_flexible",
            '{"requiere_negrita": true, "requiere_relleno": true, "requiere_bordes": true}',
            10,
            "NO"
        ),
        (
            "CRIT_06",
            "Combinar celdas del título principal en hoja Resumen",
            "Resumen",
            "A1:D1",
            "celdas_combinadas",
            '{}',
            10,
            "NO"
        ),
        (
            "CRIT_07",
            "Configurar lista desplegable de validación en celda B3",
            "Resumen",
            "B3",
            "validacion_datos",
            '{"tipo": "list"}',
            10,
            "NO"
        ),
        (
            "CRIT_08",
            "Verificar existencia de Tabla Dinámica de resumen",
            "Resumen",
            "A10",
            "tabla_dinamica",
            '{"campos_fila": ["Categoria"], "campos_valor": ["Total"]}',
            10,
            "NO"
        ),
    ]

    fill_alt = PatternFill(start_color="FFF8FAFC", end_color="FFF8FAFC", fill_type="solid")

    for r_idx, fila in enumerate(criterios_ejemplo, 2):
        for c_idx, val in enumerate(fila, 1):
            c = ws_rubrica.cell(row=r_idx, column=c_idx, value=val)
            c.border = thin_border
            if r_idx % 2 == 0:
                c.fill = fill_alt
            if c_idx in (1, 3, 4, 5, 8):
                c.alignment = align_center
            elif c_idx == 7:
                c.alignment = align_right
                c.number_format = '0.0'
            else:
                c.alignment = align_left

    # Autoajuste de anchos de columna en todas las hojas
    for ws in [ws_ventas, ws_resumen, ws_rubrica]:
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    # Guardar libro con manejo seguro de archivos abiertos en Excel
    try:
        wb.save(ruta_salida)
    except PermissionError:
        base, ext = os.path.splitext(ruta_salida)
        for i in range(1, 100):
            alt_ruta = f"{base}_{i}{ext}"
            try:
                wb.save(alt_ruta)
                ruta_salida = alt_ruta
                break
            except PermissionError:
                continue

    wb.close()
    return ruta_salida
