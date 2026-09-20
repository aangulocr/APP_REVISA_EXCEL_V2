# -*- coding: utf-8 -*-
# ============================================================================
# core/evaluators/rubric_evaluator.py — Evaluador Modo B (Smart Rubric)
# ============================================================================
"""
Motor de evaluación declarativa orientada a competencias por rúbrica (0 a 100 pts).
Aplica validaciones inteligentes y marca retroalimentación visual en el archivo
del estudiante (_rev.xlsx).
"""

import os
import time
import logging
from typing import List, Dict, Optional, Any
import openpyxl

from core.evaluators.base_evaluator import (
    BaseEvaluator,
    aplicar_marcado_error,
    aplicar_marcado_aprobado,
)
from core.models import (
    CriterioRubrica,
    ResultadoCriterio,
    ResultadoEstudiante,
    ResultadoHoja,
    ModoEvaluacion,
)
from core.rubric_parser import cargar_rubrica
from core.sheet_matcher import emparejar_hojas, es_rubrica
from core.validators import obtener_validador
from core.config import SUFIJO_REVISADO, REVISADOS_DIR
from core.com_inspector import COM_DISPONIBLE, extraer_info_pivots_com

logger = logging.getLogger(__name__)


class RubricEvaluator(BaseEvaluator):
    """Evalúa libros de estudiantes según criterios declarativos ponderados."""

    def __init__(self, criterios_predefinidos: Optional[List[CriterioRubrica]] = None):
        self.criterios_predefinidos = criterios_predefinidos

    def evaluar(
        self,
        ruta_plantilla: str,
        ruta_estudiante: str,
        carpeta_salida_revisados: str = REVISADOS_DIR
    ) -> ResultadoEstudiante:
        t_inicio = time.time()
        nombre_archivo = os.path.basename(ruta_estudiante)
        nombre_estudiante = os.path.splitext(nombre_archivo)[0]

        os.makedirs(carpeta_salida_revisados, exist_ok=True)
        ruta_salida_rev = os.path.join(
            carpeta_salida_revisados,
            f"{nombre_estudiante}{SUFIJO_REVISADO}.xlsx"
        )

        res_est = ResultadoEstudiante(
            archivo=nombre_archivo,
            nombre_estudiante=nombre_estudiante,
            modo=ModoEvaluacion.MODO_B_RUBRICA.value,
        )

        # 1. Obtener criterios de la rúbrica
        criterios = self.criterios_predefinidos
        puntos_totales_rubrica = 100.0
        if not criterios:
            criterios, puntos_totales_rubrica, advs, origen = cargar_rubrica(ruta_plantilla)
            res_est.advertencias.extend(advs)

        if not criterios:
            res_est.advertencias.append("No se encontraron criterios de rúbrica para evaluar.")
            res_est.puntos_totales = 100.0
            res_est.nota_final_100 = 0.0
            res_est.aprobado = False
            res_est.tiempo_segundos = time.time() - t_inicio
            return res_est

        res_est.puntos_totales = puntos_totales_rubrica

        # 2. Cargar libros de estudiante y plantilla
        wb_est_form = None
        wb_est_data = None
        wb_p_form = None
        wb_p_data = None

        try:
            wb_est_form = openpyxl.load_workbook(ruta_estudiante, data_only=False)
            wb_est_data = openpyxl.load_workbook(ruta_estudiante, data_only=True)
        except Exception as e:
            msg = f"Error crítico al abrir archivo de estudiante '{nombre_archivo}': {e}"
            logger.error(msg)
            res_est.advertencias.append(msg)
            res_est.nota_final_100 = 0.0
            res_est.tiempo_segundos = time.time() - t_inicio
            return res_est

        try:
            if ruta_plantilla and os.path.isfile(ruta_plantilla):
                wb_p_form = openpyxl.load_workbook(ruta_plantilla, data_only=False, read_only=True)
                wb_p_data = openpyxl.load_workbook(ruta_plantilla, data_only=True, read_only=True)
        except Exception as e:
            logger.warning(f"No se pudo cargar plantilla completa para soporte: {e}")

        # 3. Emparejamiento de hojas
        hojas_criterios = list(set(c.hoja for c in criterios if c.hoja))
        hojas_estudiante = wb_est_form.sheetnames
        mapa_hojas = emparejar_hojas(hojas_criterios, hojas_estudiante, wb_p_form, wb_est_form)

        # 4. Inspección profunda opcional vía COM para tablas dinámicas
        info_pivots_com = {}
        if COM_DISPONIBLE and any(c.tipo_validacion == "tabla_dinamica" for c in criterios):
            try:
                info_pivots_com = extraer_info_pivots_com(ruta_estudiante)
            except Exception as e:
                logger.warning(f"Inspección COM omitida: {e}")

        # 5. Evaluar cada criterio
        puntos_acumulados = 0.0
        criterios_evaluados: List[ResultadoCriterio] = []
        fallos_criterios: List[str] = []

        # Contadores por hoja
        resumen_por_hoja: Dict[str, Dict[str, Any]] = {}

        for c in criterios:
            hoja_est_nombre = None
            if c.hoja:
                info_mapa = mapa_hojas.get(c.hoja)
                if info_mapa and info_mapa.get("hoja_estudiante"):
                    hoja_est_nombre = info_mapa["hoja_estudiante"]
                elif c.hoja in wb_est_form.sheetnames:
                    hoja_est_nombre = c.hoja

            ws_e = wb_est_form[hoja_est_nombre] if hoja_est_nombre and hoja_est_nombre in wb_est_form.sheetnames else None
            ws_e_data = wb_est_data[hoja_est_nombre] if hoja_est_nombre and hoja_est_nombre in wb_est_data.sheetnames else None
            ws_p = wb_p_form[c.hoja] if wb_p_form and c.hoja in wb_p_form.sheetnames else None
            ws_p_data = wb_p_data[c.hoja] if wb_p_data and c.hoja in wb_p_data.sheetnames else None

            # Inicializar tracking de la hoja
            nombre_hoja_key = c.hoja or "General"
            if nombre_hoja_key not in resumen_por_hoja:
                resumen_por_hoja[nombre_hoja_key] = {
                    "hoja": nombre_hoja_key,
                    "hoja_estudiante": hoja_est_nombre,
                    "aciertos": 0,
                    "errores": 0,
                    "detalles": []
                }

            if c.hoja and ws_e is None:
                # Hoja no encontrada
                res_crit = ResultadoCriterio(
                    criterio_id=c.id,
                    descripcion=c.criterio,
                    aprobado=False,
                    puntos_obtenidos=0.0,
                    puntos_maximos=c.puntos,
                    mensaje_detalle=f"Hoja requerida '{c.hoja}' no encontrada en el libro.",
                    obligatorio=c.obligatorio,
                    tipo_validacion=c.tipo_validacion
                )
            else:
                # Inyectar info COM si aplica
                if c.tipo_validacion == "tabla_dinamica" and hoja_est_nombre in info_pivots_com:
                    pivots_lista = info_pivots_com[hoja_est_nombre]
                    if pivots_lista:
                        c.parametros["_com_pivot_info"] = pivots_lista[0]

                validador = obtener_validador(c.tipo_validacion)
                res_crit = validador.validar(
                    ws_estudiante=ws_e,
                    criterio=c,
                    ws_estudiante_data=ws_e_data,
                    ws_plantilla=ws_p,
                    ws_plantilla_data=ws_p_data
                )

            # Acumular resultados
            criterios_evaluados.append(res_crit)

            if res_crit.aprobado:
                puntos_acumulados += res_crit.puntos_obtenidos
                resumen_por_hoja[nombre_hoja_key]["aciertos"] += 1

                # Comentario visual positivo en la primera celda
                if ws_e and res_crit.celdas_afectadas:
                    primera = res_crit.celdas_afectadas[0]
                    try:
                        celda_obj = ws_e[primera]
                        msg_aprob = f"✅ [APROBADO: +{c.puntos:.1f} pts] {c.criterio}."
                        aplicar_marcado_aprobado(celda_obj, msg_aprob)
                    except Exception:
                        pass
            else:
                fallos_criterios.append(f"{c.id}: {c.criterio} (-{c.puntos} pts) -> {res_crit.mensaje_detalle}")
                resumen_por_hoja[nombre_hoja_key]["errores"] += 1
                resumen_por_hoja[nombre_hoja_key]["detalles"].append(f"[{c.id}] {res_crit.mensaje_detalle}")

                # Marcado visual de error con relleno rojo suave y comentario
                if ws_e and res_crit.celdas_afectadas:
                    msg_fallo = (
                        f"❌ [FALLÓ: -{c.puntos:.1f} pts] {c.criterio}.\n"
                        f"Detalle: {res_crit.mensaje_detalle}"
                    )
                    for coord_err in res_crit.celdas_afectadas:
                        try:
                            celda_obj = ws_e[coord_err]
                            aplicar_marcado_error(celda_obj, msg_fallo)
                        except Exception:
                            pass

        # 6. Calcular nota final y aprobación
        res_est.criterios_evaluados = criterios_evaluados
        res_est.criterios_fallados = fallos_criterios
        res_est.puntos_obtenidos = puntos_acumulados
        res_est.puntos_totales = puntos_totales_rubrica

        if puntos_totales_rubrica > 0:
            res_est.nota_final_100 = (puntos_acumulados / puntos_totales_rubrica) * 100.0
        else:
            res_est.nota_final_100 = 0.0

        # Penalización si falló un criterio obligatorio
        hay_obligatorio_fallado = any(
            (not r.aprobado and r.obligatorio) for r in criterios_evaluados
        )
        if hay_obligatorio_fallado:
            res_est.advertencias.append("Se reprobó al menos un criterio OBLIGATORIO.")

        res_est.aprobado = (res_est.nota_final_100 >= 70.0) and not hay_obligatorio_fallado

        # Detalle de hojas
        for _, datos_h in resumen_por_hoja.items():
            tot_h = datos_h["aciertos"] + datos_h["errores"]
            pct_h = (datos_h["aciertos"] / tot_h * 100.0) if tot_h > 0 else 0.0
            res_est.detalle_hojas.append(ResultadoHoja(
                hoja=datos_h["hoja"],
                hoja_estudiante=datos_h["hoja_estudiante"],
                aciertos=datos_h["aciertos"],
                errores=datos_h["errores"],
                porcentaje=pct_h,
                detalles=datos_h["detalles"]
            ))

        # 7. Guardar libro revisado
        try:
            wb_est_form.save(ruta_salida_rev)
            res_est.archivo_revisado = ruta_salida_rev
        except Exception as e:
            res_est.advertencias.append(f"No se pudo guardar el archivo revisado: {e}")

        # Cerrar libros
        try:
            wb_est_form.close()
            wb_est_data.close()
            if wb_p_form:
                wb_p_form.close()
            if wb_p_data:
                wb_p_data.close()
        except Exception:
            pass

        res_est.tiempo_segundos = time.time() - t_inicio
        return res_est
