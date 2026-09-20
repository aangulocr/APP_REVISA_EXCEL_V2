# -*- coding: utf-8 -*-
# ============================================================================
# core/evaluators/mirror_evaluator.py — Evaluador Modo A (Plantilla Espejo)
# ============================================================================
"""
Motor de evaluación estricta celda a celda y objeto a objeto contra plantilla
idéntica (heredado de V1) con detección heurística de desplazamientos
de filas y columnas.
"""

import os
import re
import time
import logging
from typing import List, Dict, Tuple, Optional, Any
import openpyxl
from openpyxl.utils import get_column_letter

from core.evaluators.base_evaluator import BaseEvaluator, aplicar_marcado_error
from core.models import (
    ResultadoEstudiante,
    ResultadoHoja,
    ModoEvaluacion,
)
from core.sheet_matcher import emparejar_hojas, es_rubrica
from core.shift_detector import detectar_desplazamiento
from core.config import (
    SUFIJO_REVISADO,
    REVISADOS_DIR,
    MENSAJES,
    FUNCIONES_ES_A_EN,
    FUNCIONES_AGREGACION_CONMUTATIVAS,
)
from core.com_inspector import COM_DISPONIBLE, comparar_pivots_y_graficos_com

logger = logging.getLogger(__name__)


def _normalizar_formula(formula: str) -> str:
    """Normaliza fórmulas para comparación (elimina @, _xlfn., etc.)."""
    if not formula or not isinstance(formula, str):
        return formula
    f = formula.strip().upper()
    f = f.replace("_XLFN._XLWS.", "").replace("_XLFN.", "")
    if f.startswith("=@"):
        f = "=" + f[2:]
    f = f.replace("(@", "(")
    return f


def _expandir_argumentos(args_str: str) -> List[str]:
    """Expande argumentos de funciones para comparaciones conmutativas."""
    if not args_str:
        return []
    norm_str = args_str.replace(";", ",")
    partes = [p.strip() for p in norm_str.split(",") if p.strip()]
    expandidos = []
    for parte in partes:
        p_clean = parte.replace("$", "").strip()
        if ":" in p_clean:
            try:
                min_col, min_row, max_col, max_row = openpyxl.utils.range_boundaries(p_clean)
                if min_col and min_row and max_col and max_row:
                    for c in range(min_col, max_col + 1):
                        col_let = openpyxl.utils.get_column_letter(c)
                        for r in range(min_row, max_row + 1):
                            expandidos.append(f"{col_let}{r}")
                    continue
            except Exception:
                pass
        expandidos.append(p_clean)
    return sorted(expandidos)


def _analizar_formula_agregacion(formula: str):
    if not formula or not isinstance(formula, str):
        return None, []
    f = formula.strip().upper()
    m = re.match(r"^=([A-Z0-9_\.]+)\s*\((.*)\)$", f)
    if not m:
        return None, []
    func_raw = m.group(1).strip()
    args_str = m.group(2).strip()
    if "(" in args_str or ")" in args_str:
        return None, []
    func_en = FUNCIONES_ES_A_EN.get(func_raw, func_raw)
    if func_en in FUNCIONES_AGREGACION_CONMUTATIVAS:
        celdas = _expandir_argumentos(args_str)
        return func_en, celdas
    return None, []


def _analizar_operacion_cadena(formula: str, operador: str) -> List[str]:
    if not formula or not isinstance(formula, str) or not formula.startswith("="):
        return []
    cuerpo = formula[1:].strip().upper()
    if "(" in cuerpo or ")" in cuerpo:
        return []
    otros_ops = ["-", "/", "*"] if operador == "+" else ["-", "/", "+"]
    if any(op in cuerpo for op in otros_ops):
        return []
    if operador not in cuerpo:
        return []
    partes = [p.strip().replace("$", "") for p in cuerpo.split(operador) if p.strip()]
    return sorted(partes)


def _formulas_son_equivalentes(f1_norm: str, f2_norm: str) -> bool:
    if not f1_norm or not f2_norm:
        return False
    if f1_norm == f2_norm:
        return True

    func1, celdas1 = _analizar_formula_agregacion(f1_norm)
    func2, celdas2 = _analizar_formula_agregacion(f2_norm)

    if func1 and func2 and func1 == func2 and celdas1 == celdas2:
        return True

    if func1 == "SUM":
        sumandos2 = _analizar_operacion_cadena(f2_norm, "+")
        if sumandos2 and celdas1 == sumandos2:
            return True
    if func2 == "SUM":
        sumandos1 = _analizar_operacion_cadena(f1_norm, "+")
        if sumandos1 and celdas2 == sumandos1:
            return True

    sumandos1 = _analizar_operacion_cadena(f1_norm, "+")
    sumandos2 = _analizar_operacion_cadena(f2_norm, "+")
    if sumandos1 and sumandos2 and sumandos1 == sumandos2:
        return True

    factores1 = _analizar_operacion_cadena(f1_norm, "*")
    factores2 = _analizar_operacion_cadena(f2_norm, "*")
    if factores1 and factores2 and factores1 == factores2:
        return True

    return False


def _comparar_celda_modo_a(celda_p, celda_e) -> Tuple[bool, List[str]]:
    """Compara celda de plantilla contra celda de estudiante en Modo A."""
    errores = []
    val_p = celda_p.value
    val_e = celda_e.value

    es_formula_p = isinstance(val_p, str) and val_p.strip().startswith("=")
    if es_formula_p:
        es_formula_e = isinstance(val_e, str) and val_e.strip().startswith("=")
        f_p_norm = _normalizar_formula(val_p)
        f_e_norm = _normalizar_formula(val_e) if es_formula_e else str(val_e or "")

        if f_p_norm != f_e_norm:
            if not _formulas_son_equivalentes(f_p_norm, f_e_norm):
                errores.append(f"Fórmula incorrecta. Esperada: '{val_p}' | Encontrada: '{val_e or '(vacío)'}'")
    else:
        # Si no es fórmula en plantilla, no penaliza celda en Modo A estándar (focalizado en fórmulas)
        return True, []

    return len(errores) == 0, errores


class MirrorEvaluator(BaseEvaluator):
    """Evaluador de comparación estricta celda a celda y objetos."""

    def evaluar(
        self,
        ruta_plantilla: str,
        ruta_estudiante: str,
        carpeta_salida_revisados: str = REVISADOS_DIR
    ) -> ResultadoEstudiante:
        t_inicio = time.time()
        nombre_archivo = os.path.basename(ruta_estudiante)
        nombre_estudiante = os.path.splitext(nombre_archivo)[0]

        os.makedirs(carpeta_salida_revisados, exist_ok=True)
        ruta_salida_rev = os.path.join(
            carpeta_salida_revisados,
            f"{nombre_estudiante}{SUFIJO_REVISADO}.xlsx"
        )

        res_est = ResultadoEstudiante(
            archivo=nombre_archivo,
            nombre_estudiante=nombre_estudiante,
            modo=ModoEvaluacion.MODO_A_ESPEJO.value,
        )

        try:
            wb_p = openpyxl.load_workbook(ruta_plantilla, data_only=False)
            wb_e = openpyxl.load_workbook(ruta_estudiante, data_only=False)
        except Exception as e:
            res_est.advertencias.append(f"Error al cargar libros: {e}")
            res_est.tiempo_segundos = time.time() - t_inicio
            return res_est

        hojas_p = [h for h in wb_p.sheetnames if not es_rubrica(h)]
        hojas_e = [h for h in wb_e.sheetnames if not es_rubrica(h)]
        mapa_hojas = emparejar_hojas(hojas_p, hojas_e, wb_p, wb_e)

        total_aciertos_general = 0
        total_errores_general = 0

        for nom_p in hojas_p:
            info_m = mapa_hojas.get(nom_p, {})
            nom_e = info_m.get("hoja_estudiante")

            res_hoja = ResultadoHoja(hoja=nom_p, hoja_estudiante=nom_e)

            if not nom_e or nom_e not in wb_e.sheetnames:
                res_hoja.errores = 1
                res_hoja.detalles.append(f"Hoja '{nom_p}' no encontrada en el trabajo del estudiante.")
                total_errores_general += 1
                res_est.detalle_hojas.append(res_hoja)
                continue

            ws_p = wb_p[nom_p]
            ws_e = wb_e[nom_e]

            # 1. Detector de Desplazamiento de Filas/Columnas
            offset_detectado, confianza, msg_shift = detectar_desplazamiento(ws_p, ws_e)
            if offset_detectado and msg_shift:
                res_hoja.desplazamiento = msg_shift
                res_est.advertencias.append(f"Hoja '{nom_p}': {msg_shift}")

            # 2. Comparación celda a celda
            max_r = max(ws_p.max_row or 1, ws_e.max_row or 1)
            max_c = max(ws_p.max_column or 1, ws_e.max_column or 1)

            aciertos_hoja = 0
            errores_hoja = 0

            for r in range(1, max_r + 1):
                for c in range(1, max_c + 1):
                    coord = f"{get_column_letter(c)}{r}"
                    c_p = ws_p[coord]
                    c_e = ws_e[coord]

                    if c_p.value is None and c_e.value is None:
                        continue

                    es_ok, errs = _comparar_celda_modo_a(c_p, c_e)
                    if es_ok:
                        aciertos_hoja += 1
                    else:
                        errores_hoja += len(errs)
                        msg_err = f"❌ [MODO A] {'; '.join(errs)}"
                        aplicar_marcado_error(c_e, msg_err)
                        if len(res_hoja.detalles) < 10:
                            res_hoja.detalles.append(f"Celda {coord}: {'; '.join(errs)}")

            # 3. Comparación de objetos de la hoja (Celdas combinadas)
            merged_p = {m.coord for m in ws_p.merged_cells.ranges}
            merged_e = {m.coord for m in ws_e.merged_cells.ranges}
            faltantes_merge = merged_p - merged_e
            for fm in faltantes_merge:
                errores_hoja += 1
                res_hoja.detalles.append(f"Falta combinar rango '{fm}'")

            tot_eval = aciertos_hoja + errores_hoja
            pct = (aciertos_hoja / tot_eval * 100.0) if tot_eval > 0 else 0.0

            res_hoja.aciertos = aciertos_hoja
            res_hoja.errores = errores_hoja
            res_hoja.porcentaje = pct

            total_aciertos_general += aciertos_hoja
            total_errores_general += errores_hoja
            res_est.detalle_hojas.append(res_hoja)

        # 4. Inspección profunda COM si aplica
        if COM_DISPONIBLE:
            try:
                err_com, det_com = comparar_pivots_y_graficos_com(ruta_plantilla, ruta_estudiante, mapa_hojas)
                total_errores_general += err_com
                res_est.errores_com.extend(det_com)
            except Exception as e:
                res_est.advertencias.append(f"Error en validación COM: {e}")

        # Totales
        gran_total = total_aciertos_general + total_errores_general
        res_est.total_aciertos = total_aciertos_general
        res_est.total_errores = total_errores_general
        res_est.puntos_totales = 100.0
        pct_global = (total_aciertos_general / gran_total * 100.0) if gran_total > 0 else 0.0
        res_est.puntos_obtenidos = pct_global
        res_est.nota_final_100 = pct_global
        res_est.aprobado = pct_global >= 70.0

        try:
            wb_e.save(ruta_salida_rev)
            res_est.archivo_revisado = ruta_salida_rev
        except Exception as e:
            res_est.advertencias.append(f"No se pudo guardar libro revisado: {e}")

        try:
            wb_p.close()
            wb_e.close()
        except Exception:
            pass

        res_est.tiempo_segundos = time.time() - t_inicio
        return res_est
