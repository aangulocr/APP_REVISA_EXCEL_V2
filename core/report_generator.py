# -*- coding: utf-8 -*-
# ============================================================================
# core/report_generator.py — Generador de reportes consolidados (CSV, JSON, XLSX)
# ============================================================================
"""
Genera los tres reportes oficiales de la auditoría:
  1. REPORTE_NOTAS.csv (fácil importación en sistemas escolares / SIGE)
  2. REPORTE_DETALLADO.json (trazabilidad completa y diagnósticos)
  3. LOG_NOTAS.xlsx (reporte profesional en Excel con formatos y matrices)
"""

import os
import csv
import json
import logging
from datetime import datetime
from typing import List, Dict, Any
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from core.models import ResultadoEstudiante
from core.config import (
    COLOR_HEADER_BG,
    COLOR_HEADER_FG,
    COLOR_SUCCESS_FILL,
    COLOR_ERROR_FILL,
)

logger = logging.getLogger(__name__)


def generar_observaciones_estudiante(res: ResultadoEstudiante) -> str:
    """
    Genera un texto consolidado con todos los fallos, penalizaciones y justificaciones
    de la nota obtenida por el estudiante tanto en Modo A como en Modo B.
    """
    lineas = []

    # 1. Modo B: Criterios fallados
    if res.criterios_fallados:
        for cf in res.criterios_fallados:
            lineas.append(f"• [Criterio Fallado]: {cf}")

    # 2. Modo A: Hojas con errores o penalizaciones
    if res.detalle_hojas:
        for h in res.detalle_hojas:
            pct = h.porcentaje if h.porcentaje is not None else 0.0
            if h.errores > 0 or pct < 100.0:
                if h.detalles:
                    for d in h.detalles:
                        lineas.append(f"• [{h.hoja} - {pct:.1f}%]: {d}")
                else:
                    lineas.append(f"• [{h.hoja} - {pct:.1f}%]: {h.errores} error(es) detectado(s)")
            elif h.detalles:
                # Advertencias menores aunque no baje puntos (ej: gráficos incrustados)
                for d in h.detalles:
                    if "⚠️" in d or "aviso" in d.lower() or "incrustado" in d.lower():
                        lineas.append(f"• [{h.hoja}]: {d}")

    # 3. Validación COM profunda
    if res.errores_com:
        for ec in res.errores_com:
            lineas.append(f"• [Validación COM]: {ec}")

    # 4. Advertencias globales (desplazamientos, etc.)
    if res.advertencias:
        for adv in res.advertencias:
            lineas.append(f"• [Advertencia]: {adv}")

    if not lineas:
        return "Trabajo perfecto. Cumplió todos los requisitos al 100%."

    return "\n".join(lineas)


def exportar_observaciones_txt(
    resultados: List[ResultadoEstudiante],
    ruta_txt: str,
    seccion: str = "SEC_DEFAULT"
) -> str:
    """
    Genera un archivo de texto plano con la retroalimentación individual
    de cada estudiante, ideal para copiar y pegar en Teams, Classroom, Moodle o correo.
    """
    os.makedirs(os.path.dirname(os.path.abspath(ruta_txt)), exist_ok=True)
    fecha_hoy = datetime.now().strftime("%d/%m/%Y %H:%M:%S")

    with open(ruta_txt, "w", encoding="utf-8") as f:
        f.write("=" * 80 + "\n")
        f.write("REPORTE DE OBSERVACIONES Y JUSTIFICACIÓN DE NOTAS\n")
        f.write(f"Fecha: {fecha_hoy} | Sección: {seccion} | Total Estudiantes: {len(resultados)}\n")
        f.write("=" * 80 + "\n\n")

        for res in resultados:
            estado_str = "APROBADO" if res.aprobado else "REPROBADO"
            f.write("-" * 80 + "\n")
            f.write(f"ESTUDIANTE: {res.nombre_estudiante}\n")
            f.write(f"ARCHIVO   : {res.archivo}\n")
            f.write(f"NOTA FINAL: {res.nota_final_100:.2f} / 100.0  [{estado_str}]\n")
            f.write("-" * 80 + "\n")
            f.write("OBSERVACIONES Y FALLOS DETECTADOS:\n")

            obs_txt = generar_observaciones_estudiante(res)
            for linea in obs_txt.split("\n"):
                f.write(f"  {linea}\n")

            f.write("\n")

        f.write("=" * 80 + "\n")
        f.write("FIN DEL REPORTE DE OBSERVACIONES\n")
        f.write("=" * 80 + "\n")

    return ruta_txt


def exportar_csv(
    resultados: List[ResultadoEstudiante],
    ruta_csv: str
) -> str:
    """Genera el reporte consolidado en formato CSV compatible con Excel (utf-8-sig)."""
    os.makedirs(os.path.dirname(os.path.abspath(ruta_csv)), exist_ok=True)

    encabezados = [
        "Archivo",
        "Estudiante",
        "Modo_Evaluacion",
        "Puntos_Obtenidos",
        "Puntos_Totales",
        "Nota_Final_100",
        "Estado",
        "Criterios_Fallados",
        "Advertencias",
        "Observaciones_Justificacion",
    ]

    with open(ruta_csv, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        writer.writerow(encabezados)

        for res in resultados:
            estado = "Aprobado" if res.aprobado else "Reprobado"
            crit_fallados_str = " | ".join(res.criterios_fallados) if res.criterios_fallados else "Ninguno"
            advs_str = " | ".join(res.advertencias) if res.advertencias else "Ninguna"
            obs_str = generar_observaciones_estudiante(res).replace("\n", " | ")

            writer.writerow([
                res.archivo,
                res.nombre_estudiante,
                res.modo,
                f"{res.puntos_obtenidos:.2f}",
                f"{res.puntos_totales:.2f}",
                f"{res.nota_final_100:.2f}",
                estado,
                crit_fallados_str,
                advs_str,
                obs_str,
            ])

    return ruta_csv


def exportar_json_detallado(
    resultados: List[ResultadoEstudiante],
    ruta_json: str,
    metadatos: Dict[str, Any] = None
) -> str:
    """Genera el reporte detallado estructurado en JSON."""
    os.makedirs(os.path.dirname(os.path.abspath(ruta_json)), exist_ok=True)

    data = {
        "fecha_generacion": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "metadatos": metadatos or {},
        "total_estudiantes": len(resultados),
        "total_aprobados": sum(1 for r in resultados if r.aprobado),
        "total_reprobados": sum(1 for r in resultados if not r.aprobado),
        "promedio_nota": round(sum(r.nota_final_100 for r in resultados) / len(resultados), 2) if resultados else 0.0,
        "estudiantes": [r.to_dict() for r in resultados],
    }

    with open(ruta_json, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=2, ensure_ascii=False)

    return ruta_json


def exportar_excel_log(
    resultados: List[ResultadoEstudiante],
    ruta_xlsx: str,
    seccion: str = "SEC_DEFAULT"
) -> str:
    """
    Genera el archivo LOG_NOTAS.xlsx con estética premium, resumen ejecutivo
    y desglose por criterios u hojas.
    """
    os.makedirs(os.path.dirname(os.path.abspath(ruta_xlsx)), exist_ok=True)

    wb = openpyxl.Workbook()
    ws_resumen = wb.active
    ws_resumen.title = "Resumen General"

    # Estilos
    header_fill = PatternFill(start_color="FF0056B3", end_color="FF0056B3", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")

    thin_border = Border(
        left=Side(style='thin', color='CCCCCC'),
        right=Side(style='thin', color='CCCCCC'),
        top=Side(style='thin', color='CCCCCC'),
        bottom=Side(style='thin', color='CCCCCC')
    )

    fill_aprob = PatternFill(start_color=COLOR_SUCCESS_FILL, end_color=COLOR_SUCCESS_FILL, fill_type="solid")
    fill_reprob = PatternFill(start_color=COLOR_ERROR_FILL, end_color=COLOR_ERROR_FILL, fill_type="solid")
    font_aprob = Font(name="Calibri", size=11, bold=True, color="276A3C")
    font_reprob = Font(name="Calibri", size=11, bold=True, color="9C0006")

    # 1. Encabezados de Resumen General
    encabezados = [
        "Fecha",
        "ID Estudiante",
        "Archivo",
        "Modo",
        "Puntos Obtenidos",
        "Puntos Máximos",
        "Nota Final (0-100)",
        "Estado",
        "OBSERVACIONES / JUSTIFICACIÓN DE NOTA",
        "Advertencias / Desplazamientos"
    ]

    ws_resumen.row_dimensions[1].height = 36
    for col_idx, h in enumerate(encabezados, 1):
        c = ws_resumen.cell(row=1, column=col_idx, value=h)
        c.fill = header_fill
        c.font = header_font
        c.alignment = align_center
        c.border = thin_border

    # Escribir filas de resumen
    fecha_hoy = datetime.now().strftime("%d/%m/%Y")
    for r_idx, res in enumerate(resultados, 2):
        ws_resumen.cell(row=r_idx, column=1, value=fecha_hoy).alignment = align_center
        ws_resumen.cell(row=r_idx, column=2, value=res.nombre_estudiante).alignment = align_left
        ws_resumen.cell(row=r_idx, column=3, value=res.archivo).alignment = align_left
        ws_resumen.cell(row=r_idx, column=4, value="Rúbrica (Modo B)" if "rubrica" in res.modo else "Espejo (Modo A)").alignment = align_center

        c_pts = ws_resumen.cell(row=r_idx, column=5, value=res.puntos_obtenidos)
        c_pts.alignment = align_right
        c_pts.number_format = '0.0'

        c_tot = ws_resumen.cell(row=r_idx, column=6, value=res.puntos_totales)
        c_tot.alignment = align_right
        c_tot.number_format = '0.0'

        c_nota = ws_resumen.cell(row=r_idx, column=7, value=res.nota_final_100)
        c_nota.alignment = align_right
        c_nota.number_format = '0.0'

        # Estado
        c_estado = ws_resumen.cell(row=r_idx, column=8, value="APROBADO" if res.aprobado else "REPROBADO")
        c_estado.alignment = align_center
        if res.aprobado:
            c_estado.fill = fill_aprob
            c_estado.font = font_aprob
            c_nota.font = font_aprob
        else:
            c_estado.fill = fill_reprob
            c_estado.font = font_reprob
            c_nota.font = font_reprob

        # Observaciones y Justificación
        obs_txt = generar_observaciones_estudiante(res)
        c_obs = ws_resumen.cell(row=r_idx, column=9, value=obs_txt)
        c_obs.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

        advs_txt = " | ".join(res.advertencias) if res.advertencias else ""
        c_adv = ws_resumen.cell(row=r_idx, column=10, value=advs_txt)
        c_adv.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

        for c_i in range(1, len(encabezados) + 1):
            ws_resumen.cell(row=r_idx, column=c_i).border = thin_border

    # 2. Hoja Adicional: Desglose por Hojas para Modo A
    tiene_detalle_hojas = any(len(r.detalle_hojas) > 0 for r in resultados)
    if tiene_detalle_hojas:
        ws_hojas = wb.create_sheet(title="Desglose por Hojas")
        headers_hojas = [
            "Estudiante",
            "Archivo",
            "Hoja Plantilla",
            "Hoja Estudiante",
            "Estado Hoja",
            "% Acierto",
            "Aciertos",
            "Errores",
            "Observaciones del Fallo"
        ]
        ws_hojas.row_dimensions[1].height = 36
        for c_idx, h in enumerate(headers_hojas, 1):
            c = ws_hojas.cell(row=1, column=c_idx, value=h)
            c.fill = header_fill
            c.font = header_font
            c.alignment = align_center
            c.border = thin_border

        fila_h = 2
        for res in resultados:
            for dh in res.detalle_hojas:
                ws_hojas.cell(row=fila_h, column=1, value=res.nombre_estudiante).alignment = align_left
                ws_hojas.cell(row=fila_h, column=2, value=res.archivo).alignment = align_left
                ws_hojas.cell(row=fila_h, column=3, value=dh.hoja).alignment = align_left
                ws_hojas.cell(row=fila_h, column=4, value=dh.hoja_estudiante or "(No encontrada)").alignment = align_left

                aprob_hoja = (dh.porcentaje or 0.0) >= 70.0
                c_est_h = ws_hojas.cell(row=fila_h, column=5, value="APROBADO" if aprob_hoja else "REPROBADO")
                c_est_h.alignment = align_center
                c_est_h.fill = fill_aprob if aprob_hoja else fill_reprob
                c_est_h.font = font_aprob if aprob_hoja else font_reprob

                c_pct = ws_hojas.cell(row=fila_h, column=6, value=(dh.porcentaje or 0.0) / 100.0)
                c_pct.alignment = align_right
                c_pct.number_format = '0.0%'
                c_pct.font = font_aprob if aprob_hoja else font_reprob

                c_ac = ws_hojas.cell(row=fila_h, column=7, value=dh.aciertos)
                c_ac.alignment = align_right

                c_er = ws_hojas.cell(row=fila_h, column=8, value=dh.errores)
                c_er.alignment = align_right

                det_str = " | ".join(dh.detalles) if dh.detalles else ("Hoja correcta" if (dh.porcentaje or 0.0) == 100.0 else "")
                c_det = ws_hojas.cell(row=fila_h, column=9, value=det_str)
                c_det.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

                for col_k in range(1, len(headers_hojas) + 1):
                    ws_hojas.cell(row=fila_h, column=col_k).border = thin_border

                fila_h += 1

    # 3. Hoja Adicional: Detalle por Criterio si aplica (Modo B)
    es_modo_b = any(len(r.criterios_evaluados) > 0 for r in resultados)
    if es_modo_b:
        ws_crit = wb.create_sheet(title="Matriz de Criterios")
        todos_crit_ids = []
        crit_desc_map = {}
        for r in resultados:
            for ce in r.criterios_evaluados:
                if ce.criterio_id not in todos_crit_ids:
                    todos_crit_ids.append(ce.criterio_id)
                    crit_desc_map[ce.criterio_id] = ce.descripcion

        headers_crit = ["Estudiante"] + todos_crit_ids + ["Nota Final"]
        ws_crit.row_dimensions[1].height = 36
        for c_idx, h in enumerate(headers_crit, 1):
            c = ws_crit.cell(row=1, column=c_idx, value=h)
            c.fill = header_fill
            c.font = header_font
            c.alignment = align_center
            c.border = thin_border

        for r_idx, res in enumerate(resultados, 2):
            ws_crit.cell(row=r_idx, column=1, value=res.nombre_estudiante).border = thin_border
            mapa_ce = {ce.criterio_id: ce for ce in res.criterios_evaluados}
            for c_idx, crit_id in enumerate(todos_crit_ids, 2):
                cel = ws_crit.cell(row=r_idx, column=c_idx)
                cel.border = thin_border
                cel.alignment = align_center
                if crit_id in mapa_ce:
                    item_ce = mapa_ce[crit_id]
                    cel.value = item_ce.puntos_obtenidos
                    cel.number_format = '0.0'
                    if item_ce.aprobado:
                        cel.fill = fill_aprob
                        cel.font = font_aprob
                    else:
                        cel.fill = fill_reprob
                        cel.font = font_reprob
                else:
                    cel.value = "-"

            c_nf = ws_crit.cell(row=r_idx, column=len(headers_crit), value=res.nota_final_100)
            c_nf.border = thin_border
            c_nf.alignment = align_right
            c_nf.number_format = '0.0'
            c_nf.font = font_aprob if res.aprobado else font_reprob

    # Autoajuste de anchos de columnas
    for ws in wb.worksheets:
        for col in ws.columns:
            header_val = str(ws.cell(row=1, column=col[0].column).value or "").lower()
            col_letter = get_column_letter(col[0].column)

            if "observacion" in header_val or "justificacion" in header_val:
                ws.column_dimensions[col_letter].width = 65
            elif "archivo" in header_val:
                ws.column_dimensions[col_letter].width = 32
            elif "estudiante" in header_val or "hoja" in header_val:
                ws.column_dimensions[col_letter].width = 28
            else:
                max_len = 0
                for cell in col:
                    val_str = str(cell.value or "")
                    if len(val_str) > max_len:
                        max_len = len(val_str)
                ws.column_dimensions[col_letter].width = max(min(max_len + 3, 40), 12)

    wb.save(ruta_xlsx)
    wb.close()
    return ruta_xlsx
