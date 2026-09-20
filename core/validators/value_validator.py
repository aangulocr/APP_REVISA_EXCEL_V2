# -*- coding: utf-8 -*-
# ============================================================================
# core/validators/value_validator.py — Validador de valores numéricos y texto
# ============================================================================
"""
Validador para 'valor_exacto' y 'valor_numerico'.
Comprueba el resultado calculado o el contenido de la celda admitiendo:
  - Tolerancia numérica absoluta o porcentual.
  - Normalización de texto (insensible a mayúsculas, espacios y tildes).
  - Resolución inteligente de fórmulas para archivos sin valores cacheados.
"""

import re
import math
from typing import Any, List, Set, Optional
import openpyxl.utils
from core.validators.base import BaseValidator
from core.models import CriterioRubrica, ResultadoCriterio
from core.sheet_matcher import normalizar_texto


def resolver_valor_celda(ws, ws_data, coord: str, visitados: Optional[Set[str]] = None) -> Any:
    """
    Obtiene el valor calculado de una celda.
    Si ws_data tiene el valor cacheado (de Excel), lo usa directamente.
    Si es None y ws tiene una fórmula, evalúa de forma segura expresiones
    básicas de suma, promedio y operaciones aritméticas.
    """
    if ws_data is not None:
        try:
            val_cache = ws_data[coord].value
            if val_cache is not None:
                return val_cache
        except Exception:
            pass

    if ws is None:
        return None

    try:
        val = ws[coord].value
    except Exception:
        return None

    if val is None:
        return None

    if not isinstance(val, str) or not val.strip().startswith("="):
        return val

    # Es una fórmula y no hay valor cacheado: resolución ligera
    if visitados is None:
        visitados = set()
    if coord in visitados:
        return 0.0
    visitados.add(coord)

    formula = val.strip()[1:].strip().upper()

    # 1. SUM / SUMA de rango o celdas
    m_sum = re.match(r'^(?:SUM|SUMA)\s*\(([A-Z0-9:,;\s]+)\)$', formula)
    if m_sum:
        partes = m_sum.group(1).replace(";", ",").split(",")
        total = 0.0
        for p in partes:
            p_clean = p.replace("$", "").strip()
            if ":" in p_clean:
                try:
                    min_c, min_r, max_c, max_r = openpyxl.utils.range_boundaries(p_clean)
                    for r in range(min_r, max_r + 1):
                        for c in range(min_c, max_c + 1):
                            let = openpyxl.utils.get_column_letter(c)
                            v = resolver_valor_celda(ws, ws_data, f"{let}{r}", visitados.copy())
                            if isinstance(v, (int, float)):
                                total += float(v)
                except Exception:
                    pass
            else:
                v = resolver_valor_celda(ws, ws_data, p_clean, visitados.copy())
                if isinstance(v, (int, float)):
                    total += float(v)
        return total

    # 2. AVERAGE / PROMEDIO
    m_avg = re.match(r'^(?:AVERAGE|PROMEDIO)\s*\(([A-Z0-9:,;\s]+)\)$', formula)
    if m_avg:
        partes = m_avg.group(1).replace(";", ",").split(",")
        vals = []
        for p in partes:
            p_clean = p.replace("$", "").strip()
            if ":" in p_clean:
                try:
                    min_c, min_r, max_c, max_r = openpyxl.utils.range_boundaries(p_clean)
                    for r in range(min_r, max_r + 1):
                        for c in range(min_c, max_c + 1):
                            let = openpyxl.utils.get_column_letter(c)
                            v = resolver_valor_celda(ws, ws_data, f"{let}{r}", visitados.copy())
                            if isinstance(v, (int, float)):
                                vals.append(float(v))
                except Exception:
                    pass
            else:
                v = resolver_valor_celda(ws, ws_data, p_clean, visitados.copy())
                if isinstance(v, (int, float)):
                    vals.append(float(v))
        return (sum(vals) / len(vals)) if vals else 0.0

    # 3. Operaciones aritméticas simples: C2*D2, E2-F2, E2*0.05
    def repl_coord(m):
        c_ref = m.group(0)
        v = resolver_valor_celda(ws, ws_data, c_ref, visitados.copy())
        if isinstance(v, (int, float)):
            return str(v)
        return "0"

    expr = re.sub(r'\$?[A-Z]+\$?[0-9]+', repl_coord, formula)
    # Reemplazar porcentajes como 5% por 0.05
    expr = re.sub(r'(\d+(?:\.\d+)?)%', r'(\1/100.0)', expr)
    try:
        # Solo permitir operaciones matemáticas seguras
        if re.match(r'^[0-9\.\s\+\-\*\/\(\)]+$', expr):
            return eval(expr, {"__builtins__": None}, {})
    except Exception:
        pass

    return val


class ValueValidator(BaseValidator):
    """Evalúa coincidencia de valor exacto o numérico con tolerancias."""

    def validar(
        self,
        ws_estudiante,
        criterio: CriterioRubrica,
        ws_estudiante_data=None,
        ws_plantilla=None,
        ws_plantilla_data=None
    ) -> ResultadoCriterio:
        params = criterio.parametros or {}
        coords = self.obtener_coordenadas_rango(criterio.rango)

        if not coords:
            return ResultadoCriterio(
                criterio_id=criterio.id,
                descripcion=criterio.criterio,
                aprobado=False,
                puntos_obtenidos=0.0,
                puntos_maximos=criterio.puntos,
                mensaje_detalle=f"Rango '{criterio.rango}' no especificado o inválido.",
                obligatorio=criterio.obligatorio,
                tipo_validacion=criterio.tipo_validacion
            )

        tolerancia_abs = float(params.get("tolerancia", 0.01))
        tolerancia_pct = float(params.get("tolerancia_pct", 0.0))

        valores_esperados = params.get("valores")
        valor_esperado_unico = params.get("valor")

        fallos: List[str] = []
        celdas_erroneas: List[str] = []

        for idx, coord in enumerate(coords):
            # Obtener valor estudiante resolviendo fórmula si fuera necesario
            val_est = resolver_valor_celda(ws_estudiante, ws_estudiante_data, coord)

            # Determinar valor esperado para esta celda
            val_esp = None
            if valores_esperados and isinstance(valores_esperados, list) and idx < len(valores_esperados):
                val_esp = valores_esperados[idx]
            elif valor_esperado_unico is not None:
                val_esp = valor_esperado_unico
            else:
                val_esp = resolver_valor_celda(ws_plantilla, ws_plantilla_data, coord)

            coincide, motivo = self._comparar_un_valor(val_est, val_esp, tolerancia_abs, tolerancia_pct)
            if not coincide:
                celdas_erroneas.append(coord)
                fallos.append(f"Celda {coord}: {motivo}")

        aprobado = len(fallos) == 0
        puntos = criterio.puntos if aprobado else 0.0
        mensaje = "Valor correcto." if aprobado else "; ".join(fallos[:3])
        if len(fallos) > 3:
            mensaje += f" (y {len(fallos)-3} celdas más con diferencias)"

        return ResultadoCriterio(
            criterio_id=criterio.id,
            descripcion=criterio.criterio,
            aprobado=aprobado,
            puntos_obtenidos=puntos,
            puntos_maximos=criterio.puntos,
            mensaje_detalle=mensaje,
            celdas_afectadas=celdas_erroneas if not aprobado else coords[:1],
            obligatorio=criterio.obligatorio,
            tipo_validacion=criterio.tipo_validacion
        )

    def _comparar_un_valor(
        self,
        val_est: Any,
        val_esp: Any,
        tol_abs: float,
        tol_pct: float
    ) -> tuple[bool, str]:
        if val_esp is None and val_est is None:
            return True, ""
        if val_est is None:
            return False, f"Celda vacía (se esperaba '{val_esp}')"
        if val_esp is None:
            return False, f"Se esperaba celda vacía pero contiene '{val_est}'"

        # Comparación numérica
        es_num_est, num_est = self._a_numero(val_est)
        es_num_esp, num_esp = self._a_numero(val_esp)

        if es_num_est and es_num_esp:
            dif = abs(num_est - num_esp)
            if dif <= tol_abs:
                return True, ""
            if tol_pct > 0 and num_esp != 0:
                if (dif / abs(num_esp)) <= tol_pct:
                    return True, ""
            return False, f"Esperaba {num_esp}, obtenido {num_est} (dif: {dif:.4f})"

        # Comparación de texto flexible
        txt_est = normalizar_texto(str(val_est))
        txt_esp = normalizar_texto(str(val_esp))

        if txt_est == txt_esp:
            return True, ""

        return False, f"Esperaba '{val_esp}', obtenido '{val_est}'"

    @staticmethod
    def _a_numero(val: Any) -> tuple[bool, float]:
        if isinstance(val, (int, float)) and not isinstance(val, bool):
            if math.isnan(val) or math.isinf(val):
                return False, 0.0
            return True, float(val)
        if isinstance(val, str):
            v_clean = val.strip().replace("$", "").replace("€", "").replace("₡", "").replace("%", "").strip()
            if "," in v_clean and "." not in v_clean:
                v_clean = v_clean.replace(",", ".")
            try:
                return True, float(v_clean)
            except ValueError:
                return False, 0.0
        return False, 0.0
